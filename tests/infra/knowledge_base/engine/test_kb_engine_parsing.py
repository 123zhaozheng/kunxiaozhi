"""Document parsing for the personal knowledge base (design §5)."""

from __future__ import annotations

import io
import json
import zipfile
from types import SimpleNamespace
from typing import Any

import docx
import openpyxl
import pytest
from pypdf import PdfWriter

from src.infra.knowledge_base import parsing
from src.infra.knowledge_base.parsing import (
    ParseError,
    is_supported_filename,
    mineru_available,
    normalize_extension,
    parse_document,
)

MINERU_ON = SimpleNamespace(
    ENABLE_DOCUMENT_PARSE=True, MINERU_API_BASE_URL="http://mineru:8000", MINERU_API_KEY="secret"
)
MINERU_OFF = SimpleNamespace(
    ENABLE_DOCUMENT_PARSE=False, MINERU_API_BASE_URL="http://mineru:8000", MINERU_API_KEY=""
)


class FakeMinerU:
    """Replaces MinerUClient; class-level state is reset by the fixture."""

    calls: list[dict[str, Any]] = []
    result: str | Exception = "# parsed by mineru\n\nbody"

    def __init__(self, base_url: str, api_key: str | None = None) -> None:
        self.base_url = base_url
        self.api_key = api_key

    async def parse_bytes(self, filename: str, content: bytes, content_type: str) -> str:
        FakeMinerU.calls.append(
            {
                "base_url": self.base_url,
                "api_key": self.api_key,
                "filename": filename,
                "content_type": content_type,
                "size": len(content),
            }
        )
        if isinstance(FakeMinerU.result, Exception):
            raise FakeMinerU.result
        return FakeMinerU.result


@pytest.fixture
def mineru(monkeypatch: pytest.MonkeyPatch) -> type[FakeMinerU]:
    FakeMinerU.calls = []
    FakeMinerU.result = "# parsed by mineru\n\nbody"
    monkeypatch.setattr(parsing, "MinerUClient", FakeMinerU)
    monkeypatch.setattr(parsing, "settings", MINERU_ON)
    return FakeMinerU


@pytest.fixture
def no_mineru(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(parsing, "settings", MINERU_OFF)


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("项目计划", level=1)
    document.add_paragraph("这是第一段。")
    document.add_heading("子标题", level=2)
    document.add_paragraph("要点一", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "值|x"
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "销售"
    sheet.append(["产品", "数量", "备注"])
    sheet.append(["苹果", 3, None])
    sheet.append([None, None, None])
    sheet.append(["梨", 4.5, "多\n行"])
    lone = workbook.create_sheet("单行")
    lone.append(["只有表头", "B"])
    workbook.create_sheet("空表")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_text_pdf(text: str) -> bytes:
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        None,
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    stream = f"BT /F1 14 Tf 20 100 Td ({text}) Tj ET".encode()
    objects[3] = b"<</Length %d>>\nstream\n%s\nendstream" % (len(stream), stream)
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<</Root 1 0 R/Size %d>>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_at,
    )
    return bytes(out)


def make_zip(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return buffer.getvalue()


# --- helpers -----------------------------------------------------------------------


def test_normalize_extension_and_support() -> None:
    assert normalize_extension("Report.PDF") == ".pdf"
    assert normalize_extension("a/b\\c.Docx") == ".docx"
    assert normalize_extension("archive.tar.gz") == ".gz"
    assert normalize_extension("noext") == ""
    assert is_supported_filename("a.MD")
    assert is_supported_filename("scan.TIFF")
    assert not is_supported_filename("a.exe")
    assert not is_supported_filename("a.doc")
    assert not is_supported_filename("README")


@pytest.mark.parametrize(
    ("enabled", "base_url", "expected"),
    [(True, "http://m", True), (True, "  ", False), (True, "", False), (False, "http://m", False)],
)
def test_mineru_available(
    monkeypatch: pytest.MonkeyPatch, enabled: bool, base_url: str, expected: bool
) -> None:
    monkeypatch.setattr(
        parsing,
        "settings",
        SimpleNamespace(ENABLE_DOCUMENT_PARSE=enabled, MINERU_API_BASE_URL=base_url),
    )
    assert mineru_available() is expected


async def test_rejects_unsupported_and_empty(no_mineru: None) -> None:
    with pytest.raises(ParseError, match="不支持的文件类型"):
        await parse_document("evil.exe", b"MZ")
    with pytest.raises(ParseError, match="文档内容为空"):
        await parse_document("a.txt", b"")
    with pytest.raises(ParseError, match="文档内容为空"):
        await parse_document("a.txt", b"  \n\t \n")


# --- text-like formats -------------------------------------------------------------


async def test_txt_and_md_decode_utf8_bom_and_crlf(no_mineru: None) -> None:
    result = await parse_document("a.txt", "第一行\r\n第二行\r".encode("utf-8-sig"))
    assert result.method == "local"
    assert result.markdown == "第一行\n第二行"
    md = await parse_document("notes.md", b"# Title\n\ntext")
    assert md.markdown == "# Title\n\ntext"


async def test_txt_detects_gbk(no_mineru: None) -> None:
    text = "知识库是企业沉淀经验的重要工具。员工可以在其中上传文档、检索资料并分享给同事，提升协作效率。"
    result = await parse_document("a.txt", text.encode("gbk"))
    assert result.markdown == text


async def test_txt_binary_garbage_without_encoding_raises(no_mineru: None) -> None:
    with pytest.raises(ParseError):
        await parse_document("a.txt", bytes(range(0, 256)) * 4)


async def test_html_to_markdown_strips_script_and_style(no_mineru: None) -> None:
    html = (
        b"<html><head><title>T</title><style>p{color:red}</style></head><body>"
        b"<h1>Hello</h1><script>alert('x')</script><p>world <b>bold</b></p></body></html>"
    )
    result = await parse_document("page.htm", html)
    assert "# Hello" in result.markdown
    assert "**bold**" in result.markdown
    assert "alert" not in result.markdown
    assert "color:red" not in result.markdown


async def test_json_is_pretty_printed_in_fence(no_mineru: None) -> None:
    result = await parse_document("a.json", json.dumps({"名称": "值", "n": [1, 2]}).encode())
    assert result.markdown.startswith("```json\n{")
    assert result.markdown.endswith("}\n```")
    assert '"名称": "值"' in result.markdown


async def test_invalid_json_falls_back_to_plain_text(no_mineru: None) -> None:
    result = await parse_document("a.json", b"{not json at all")
    assert result.markdown == "{not json at all"


async def test_csv_rows_become_header_value_lines(no_mineru: None) -> None:
    csv_bytes = '产品,数量,备注\n苹果,3,\n,,\n梨,4,"a,b"\n多余,1,x,y\n'.encode()
    result = await parse_document("a.csv", csv_bytes)
    assert result.markdown.splitlines() == [
        "产品: 苹果 | 数量: 3",
        "产品: 梨 | 数量: 4 | 备注: a,b",
        "产品: 多余 | 数量: 1 | 备注: x | col4: y",
    ]


async def test_csv_detects_semicolon_and_tab(no_mineru: None) -> None:
    semi = await parse_document("a.csv", b"a;b\n1;2\n")
    assert semi.markdown == "a: 1 | b: 2"
    tab = await parse_document("a.csv", b"a\tb\n1\t2\n")
    assert tab.markdown == "a: 1 | b: 2"


async def test_csv_header_only_keeps_header(no_mineru: None) -> None:
    result = await parse_document("a.csv", b"name,age\n")
    assert result.markdown == "name | age"


async def test_xlsx_sheets_are_sections(no_mineru: None) -> None:
    result = await parse_document("a.xlsx", make_xlsx())
    assert result.method == "local"
    assert result.markdown == (
        "## 销售\n产品: 苹果 | 数量: 3\n产品: 梨 | 数量: 4.5 | 备注: 多 行\n\n## 单行\n只有表头 | B"
    )


async def test_xlsx_corrupt_raises_parse_error(no_mineru: None) -> None:
    with pytest.raises(ParseError, match="损坏"):
        await parse_document("a.xlsx", b"definitely not a workbook")


class _FakeCell:
    def __init__(self, ctype: int, value: Any) -> None:
        self.ctype = ctype
        self.value = value


class _FakeSheet:
    def __init__(self, name: str, rows: list[list[_FakeCell]]) -> None:
        self.name = name
        self._rows = rows
        self.nrows = len(rows)

    def row(self, index: int) -> list[_FakeCell]:
        return self._rows[index]


class _FakeBook:
    datemode = 0

    def __init__(self, sheets: list[_FakeSheet]) -> None:
        self._sheets = sheets
        self.released = False

    def sheets(self) -> list[_FakeSheet]:
        return self._sheets

    def release_resources(self) -> None:
        self.released = True


async def test_xls_rows_via_xlrd(monkeypatch: pytest.MonkeyPatch, no_mineru: None) -> None:
    import xlrd

    text, number, empty = xlrd.XL_CELL_TEXT, xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_EMPTY
    date = xlrd.XL_CELL_DATE
    book = _FakeBook(
        [
            _FakeSheet(
                "Sheet1",
                [
                    [_FakeCell(text, "名称"), _FakeCell(text, "数量"), _FakeCell(text, "日期")],
                    [_FakeCell(text, "苹果"), _FakeCell(number, 3.0), _FakeCell(date, 45000.0)],
                    [_FakeCell(empty, ""), _FakeCell(empty, ""), _FakeCell(empty, "")],
                ],
            )
        ]
    )
    monkeypatch.setattr(xlrd, "open_workbook", lambda **kwargs: book)
    result = await parse_document("a.xls", b"fake-xls-bytes")
    assert result.markdown == "## Sheet1\n名称: 苹果 | 数量: 3 | 日期: 2023-03-15"
    assert book.released


async def test_xls_invalid_bytes_raise_parse_error(no_mineru: None) -> None:
    with pytest.raises(ParseError, match="损坏"):
        await parse_document("a.xls", b"not an xls file")


# --- zip ---------------------------------------------------------------------------


async def test_zip_prefers_full_md(no_mineru: None) -> None:
    archive = make_zip(
        {
            "docs/other.md": b"other",
            "docs/FULL.md": "# 全文\n![](images/a.jpg)\n正文".encode(),
            "images/a.jpg": b"\xff\xd8",
        }
    )
    result = await parse_document("bundle.zip", archive)
    assert result.markdown == "# 全文\n正文"


async def test_zip_falls_back_to_first_md_and_skips_macos_entries(no_mineru: None) -> None:
    archive = make_zip(
        {"__MACOSX/._a.md": b"junk", "readme.txt": b"no", "a.md": b"first", "b.md": b"second"}
    )
    result = await parse_document("bundle.zip", archive)
    assert result.markdown == "first"


async def test_zip_without_markdown_raises(no_mineru: None) -> None:
    with pytest.raises(ParseError, match="压缩包中未找到 .md 文件"):
        await parse_document("bundle.zip", make_zip({"a.txt": b"text"}))


@pytest.mark.parametrize("name", ["../evil.md", "/abs/evil.md", "a/../../evil.md", "C:/evil.md"])
async def test_zip_rejects_unsafe_paths(no_mineru: None, name: str) -> None:
    with pytest.raises(ParseError, match="不安全"):
        await parse_document("bundle.zip", make_zip({name: b"x", "ok.md": b"ok"}))


async def test_zip_rejects_too_many_members(no_mineru: None) -> None:
    members = {f"f{i}.txt": b"x" for i in range(1001)}
    with pytest.raises(ParseError, match="数量过多"):
        await parse_document("bundle.zip", make_zip(members))


async def test_zip_rejects_oversized_member(
    monkeypatch: pytest.MonkeyPatch, no_mineru: None
) -> None:
    monkeypatch.setattr(parsing, "_ZIP_MAX_BYTES", 10)
    with pytest.raises(ParseError, match="过大"):
        await parse_document("bundle.zip", make_zip({"a.md": b"x" * 100}))


async def test_zip_corrupt_raises(no_mineru: None) -> None:
    with pytest.raises(ParseError, match="损坏"):
        await parse_document("bundle.zip", b"PK-not-really")


# --- docx / pdf --------------------------------------------------------------------


async def test_docx_local_fallback_headings_lists_tables(no_mineru: None) -> None:
    result = await parse_document("plan.docx", make_docx())
    assert result.method == "local"
    assert result.markdown == (
        "# 项目计划\n\n这是第一段。\n\n## 子标题\n\n- 要点一\n\n"
        "| 名称 | 值\\|x |\n| --- | --- |\n| a | b |"
    )


async def test_docx_prefers_mineru_when_available(mineru: type[FakeMinerU]) -> None:
    result = await parse_document("plan.docx", make_docx())
    assert result.method == "mineru"
    assert result.markdown == "# parsed by mineru\n\nbody"
    (call,) = mineru.calls
    assert call["base_url"] == "http://mineru:8000"
    assert call["api_key"] == "secret"
    assert call["filename"] == "plan.docx"
    assert call["content_type"].endswith("wordprocessingml.document")


async def test_docx_falls_back_when_mineru_fails(mineru: type[FakeMinerU]) -> None:
    mineru.result = RuntimeError("boom")
    result = await parse_document("plan.docx", make_docx())
    assert result.method == "local"
    assert result.markdown.startswith("# 项目计划")
    assert len(mineru.calls) == 1


async def test_docx_corrupt_without_mineru_raises(no_mineru: None) -> None:
    with pytest.raises(ParseError, match="损坏"):
        await parse_document("plan.docx", b"not a docx")


async def test_pdf_text_layer_when_mineru_unavailable(no_mineru: None) -> None:
    result = await parse_document("a.pdf", make_text_pdf("Hello KB World"))
    assert result.method == "pdf_text"
    assert result.markdown == "Hello KB World"


async def test_pdf_mineru_success_and_fallback(mineru: type[FakeMinerU]) -> None:
    pdf = make_text_pdf("Hello KB World")
    ok = await parse_document("a.pdf", pdf)
    assert ok.method == "mineru"
    assert mineru.calls[0]["content_type"] == "application/pdf"

    mineru.result = RuntimeError("down")
    fallback = await parse_document("a.pdf", pdf)
    assert (fallback.method, fallback.markdown) == ("pdf_text", "Hello KB World")

    mineru.result = "   \n"
    blank = await parse_document("a.pdf", pdf)
    assert blank.method == "pdf_text"


async def test_pdf_without_text_layer_raises_admin_hint(no_mineru: None) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ParseError, match="PDF 没有可提取的文本层，需要管理员启用文档解析服务"):
        await parse_document("scan.pdf", buffer.getvalue())


async def test_pdf_corrupt_raises(no_mineru: None) -> None:
    with pytest.raises(ParseError, match="损坏"):
        await parse_document("a.pdf", b"%PDF-1.4 garbage")


# --- mineru-only formats -----------------------------------------------------------


@pytest.mark.parametrize(
    "name", ["deck.pptx", "p.png", "p.JPG", "p.jpeg", "p.webp", "p.bmp", "p.tif"]
)
async def test_mineru_only_formats_require_service(no_mineru: None, name: str) -> None:
    with pytest.raises(ParseError, match="需要管理员启用文档解析服务"):
        await parse_document(name, b"bytes")


@pytest.mark.parametrize(
    ("name", "content_type"),
    [
        ("deck.pptx", "presentationml.presentation"),
        ("p.png", "image/png"),
        ("p.jpg", "image/jpeg"),
        ("p.tiff", "image/tiff"),
        ("p.bmp", "image/bmp"),
    ],
)
async def test_mineru_only_formats_use_mineru(
    mineru: type[FakeMinerU], name: str, content_type: str
) -> None:
    result = await parse_document(name, b"bytes", "application/octet-stream")
    assert result.method == "mineru"
    assert content_type in mineru.calls[0]["content_type"]


async def test_mineru_only_failure_is_a_parse_error(mineru: type[FakeMinerU]) -> None:
    mineru.result = RuntimeError("http://internal-host:8000 refused")
    with pytest.raises(ParseError) as excinfo:
        await parse_document("deck.pptx", b"bytes")
    assert "internal-host" not in str(excinfo.value)


async def test_mineru_output_is_normalized(mineru: type[FakeMinerU]) -> None:
    mineru.result = "\r\n# T\r\n\r\ntext\x00\r\n"
    result = await parse_document("deck.pptx", b"bytes")
    assert result.markdown == "# T\n\ntext"
