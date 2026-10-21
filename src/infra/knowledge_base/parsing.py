"""Document parsing: file bytes -> markdown (design §5)."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import PurePosixPath
from typing import Any, Literal

from charset_normalizer import from_bytes

from src.infra.async_utils import run_blocking_io
from src.infra.logging import get_logger
from src.infra.tool.mineru_client import MinerUClient, strip_server_local_image_links
from src.kernel.config import settings

logger = get_logger(__name__)

ParseMethod = Literal["mineru", "pdf_text", "local"]

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(
    {
        ".pdf",
        ".docx",
        ".pptx",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".bmp",
        ".tif",
        ".tiff",
        ".txt",
        ".md",
        ".html",
        ".htm",
        ".json",
        ".csv",
        ".xlsx",
        ".xls",
        ".zip",
    }
)
# Extensions that can only be parsed by MinerU (no local fallback).
MINERU_ONLY_EXTENSIONS: frozenset[str] = frozenset(
    {".pptx", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
)


class ParseError(Exception):
    """User-presentable parse failure (message is Chinese, shown in the UI)."""


@dataclass(frozen=True)
class ParsedDocument:
    markdown: str
    method: ParseMethod


_MINERU_CONTENT_TYPES: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}
_MINERU_REQUIRED_MESSAGE = "需要管理员启用文档解析服务"
_PDF_NO_TEXT_MESSAGE = "PDF 没有可提取的文本层，需要管理员启用文档解析服务"
_EMPTY_MESSAGE = "文档内容为空"
_ZIP_MAX_MEMBERS = 1000
_ZIP_MAX_BYTES = 64 * 1024 * 1024
_ROW_SEPARATOR = " | "


def normalize_extension(filename: str) -> str:
    """Return the lower-cased final suffix including the dot ('' when none)."""
    return PurePosixPath((filename or "").replace("\\", "/")).suffix.lower()


def is_supported_filename(filename: str) -> bool:
    return normalize_extension(filename) in SUPPORTED_EXTENSIONS


def mineru_available() -> bool:
    """True iff settings.ENABLE_DOCUMENT_PARSE and MINERU_API_BASE_URL is non-empty."""
    return bool(getattr(settings, "ENABLE_DOCUMENT_PARSE", False)) and bool(
        str(getattr(settings, "MINERU_API_BASE_URL", "") or "").strip()
    )


# --- text helpers -------------------------------------------------------------------


def _decode_bytes(data: bytes) -> str:
    # Valid UTF-8 is by far the common case and cheaper/more reliable than detection.
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    match = from_bytes(data).best()
    if match is None:
        raise ParseError("无法识别文件编码")
    return str(match)


def _normalize_markdown(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    return text.strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return (
            value.isoformat(sep=" ") if value.time() != datetime.min.time() else str(value.date())
        )
    if isinstance(value, date):
        return value.isoformat()
    return re.sub(r"\s*[\r\n]+\s*", " ", str(value)).strip()


def _rows_to_lines(rows: Iterator[list[str]]) -> list[str]:
    """First non-empty row is the header; later rows become 'h1: v1 | h2: v2' lines."""
    header: list[str] | None = None
    lines: list[str] = []
    for row in rows:
        if not any(row):
            continue
        if header is None:
            header = row
            continue
        parts = []
        for position, value in enumerate(row):
            if not value:
                continue
            label = header[position] if position < len(header) and header[position] else None
            parts.append(f"{label or f'col{position + 1}'}: {value}")
        if parts:
            lines.append(_ROW_SEPARATOR.join(parts))
    if not lines and header is not None:
        # A single-row sheet has no data rows; keep its content instead of dropping it.
        lines.append(_ROW_SEPARATOR.join(cell for cell in header if cell))
    return lines


# --- local parsers (blocking) -------------------------------------------------------


def _parse_text(data: bytes) -> str:
    return _decode_bytes(data)


def _parse_html(data: bytes) -> str:
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    soup = BeautifulSoup(_decode_bytes(data), "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    markdown = markdownify(str(soup), heading_style="ATX")
    return re.sub(r"\n{3,}", "\n\n", markdown)


def _parse_json(data: bytes) -> str:
    text = _decode_bytes(data)
    try:
        pretty = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except (ValueError, RecursionError):
        return text
    return f"```json\n{pretty}\n```"


def _parse_csv(data: bytes) -> str:
    text = _decode_bytes(data)
    first_line = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(",\t;|", key=lambda candidate: (first_line.count(candidate), candidate == ","))
    try:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        return "\n".join(_rows_to_lines([_cell_text(cell) for cell in row] for row in reader))
    except csv.Error as exc:
        raise ParseError("CSV 文件格式无法解析") from exc


def _parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sections: list[str] = []
        for sheet in workbook.worksheets:
            rows = ([_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True))
            lines = _rows_to_lines(rows)
            if lines:
                sections.append("\n".join([f"## {sheet.title}", *lines]))
        return "\n\n".join(sections)
    finally:
        workbook.close()


def _parse_xls(data: bytes) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:

        def convert(cell: Any) -> str:
            if cell.ctype == xlrd.XL_CELL_DATE:
                return _cell_text(xlrd.xldate_as_datetime(cell.value, book.datemode))
            if cell.ctype == xlrd.XL_CELL_BOOLEAN:
                return _cell_text(bool(cell.value))
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_ERROR, xlrd.XL_CELL_BLANK):
                return ""
            return _cell_text(cell.value)

        sections: list[str] = []
        for sheet in book.sheets():
            rows = ([convert(cell) for cell in sheet.row(i)] for i in range(sheet.nrows))
            lines = _rows_to_lines(rows)
            if lines:
                sections.append("\n".join([f"## {sheet.name}", *lines]))
        return "\n\n".join(sections)
    finally:
        book.release_resources()


def _is_unsafe_zip_path(name: str) -> bool:
    normalized = name.replace("\\", "/")
    return (
        normalized.startswith("/")
        or re.match(r"^[A-Za-z]:", normalized) is not None
        or ".." in PurePosixPath(normalized).parts
    )


def _parse_zip(data: bytes) -> str:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ParseError("压缩包已损坏或格式不受支持") from exc
    with archive:
        members = archive.infolist()
        if len(members) > _ZIP_MAX_MEMBERS:
            raise ParseError("压缩包内文件数量过多")
        if any(_is_unsafe_zip_path(member.filename) for member in members):
            raise ParseError("压缩包包含不安全的路径")
        candidates = [
            member
            for member in members
            if not member.is_dir()
            and member.filename.lower().endswith(".md")
            and "__MACOSX/" not in member.filename
            and not PurePosixPath(member.filename).name.startswith("._")
        ]
        chosen = next(
            (m for m in candidates if PurePosixPath(m.filename).name.lower() == "full.md"),
            candidates[0] if candidates else None,
        )
        if chosen is None:
            raise ParseError("压缩包中未找到 .md 文件")
        if chosen.file_size > _ZIP_MAX_BYTES:
            raise ParseError("压缩包内文件解压后过大")
        with archive.open(chosen) as handle:
            content = handle.read(_ZIP_MAX_BYTES + 1)
        if len(content) > _ZIP_MAX_BYTES:
            raise ParseError("压缩包内文件解压后过大")
    return strip_server_local_image_links(_decode_bytes(content))


def _markdown_table(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    padded = [[*row, *[""] * (width - len(row))] for row in rows]
    escaped = [[cell.replace("|", "\\|") for cell in row] for row in padded]
    lines = ["| " + " | ".join(escaped[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in escaped[1:])
    return "\n".join(lines)


def _heading_prefix(style_name: str) -> str:
    name = style_name.strip()
    if name.lower() == "title":
        return "# "
    match = re.match(r"^(?:heading|标题)\s*(\d)$", name, re.IGNORECASE)
    if match and 1 <= int(match.group(1)) <= 6:
        return "#" * int(match.group(1)) + " "
    if name.lower().startswith("list bullet"):
        return "- "
    return ""


def _parse_docx_local(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    blocks: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if text:
                style = item.style.name if item.style is not None and item.style.name else ""
                blocks.append(_heading_prefix(style) + text)
        elif isinstance(item, Table):
            rows = [
                [re.sub(r"\s*[\r\n]+\s*", " ", cell.text).strip() for cell in row.cells]
                for row in item.rows
            ]
            rows = [row for row in rows if any(row)]
            if rows:
                blocks.append(_markdown_table(rows))
    return "\n\n".join(blocks)


def _parse_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ParseError("PDF 已加密，无法提取文本")
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


_LOCAL_PARSERS: dict[str, Callable[[bytes], str]] = {
    ".txt": _parse_text,
    ".md": _parse_text,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".json": _parse_json,
    ".csv": _parse_csv,
    ".xlsx": _parse_xlsx,
    ".xls": _parse_xls,
    ".zip": _parse_zip,
}


async def _run_local(parser: Callable[[bytes], str], data: bytes, ext: str) -> str:
    try:
        return await run_blocking_io(parser, data)
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001 - third-party parsers raise many types
        logger.warning(f"[KB] Local parse of {ext} failed: {type(exc).__name__}: {exc}")
        raise ParseError("文件已损坏或格式不受支持") from exc


async def _parse_with_mineru(filename: str, data: bytes, ext: str, content_type: str | None) -> str:
    client = MinerUClient(
        base_url=str(settings.MINERU_API_BASE_URL),
        api_key=getattr(settings, "MINERU_API_KEY", "") or None,
    )
    markdown = await client.parse_bytes(
        filename, data, _MINERU_CONTENT_TYPES.get(ext) or content_type or "application/octet-stream"
    )
    return markdown


async def _try_mineru(filename: str, data: bytes, ext: str, content_type: str | None) -> str | None:
    """MinerU markdown, or None when it fails or returns nothing."""
    try:
        markdown = _normalize_markdown(await _parse_with_mineru(filename, data, ext, content_type))
    except Exception as exc:  # noqa: BLE001 - callers fall back or raise a friendly error
        logger.warning(f"[KB] MinerU failed for {ext} file: {type(exc).__name__}: {exc}")
        return None
    return markdown or None


async def parse_document(
    filename: str, data: bytes, content_type: str | None = None
) -> ParsedDocument:
    """Parse raw bytes into markdown.

    Raises ParseError for unsupported/empty/unparseable input. Blocking parsers run via
    run_blocking_io. Never returns empty markdown (raises ParseError instead).
    """
    ext = normalize_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError(f"不支持的文件类型：{ext or filename}")
    if not data:
        raise ParseError(_EMPTY_MESSAGE)

    markdown: str | None = None
    method: ParseMethod = "local"
    use_mineru = mineru_available()

    if ext in MINERU_ONLY_EXTENSIONS:
        if not use_mineru:
            raise ParseError(_MINERU_REQUIRED_MESSAGE)
        markdown = await _try_mineru(filename, data, ext, content_type)
        if markdown is None:
            raise ParseError("文档解析服务暂时不可用或未能识别内容，请稍后重试")
        method = "mineru"
    elif ext in (".pdf", ".docx"):
        if use_mineru:
            markdown = await _try_mineru(filename, data, ext, content_type)
            method = "mineru"
        if markdown is None:
            if ext == ".pdf":
                text = await _run_local(_parse_pdf_text, data, ext)
                if not text.strip():
                    raise ParseError(_PDF_NO_TEXT_MESSAGE)
                markdown, method = _normalize_markdown(text), "pdf_text"
            else:
                markdown, method = (
                    _normalize_markdown(await _run_local(_parse_docx_local, data, ext)),
                    "local",
                )
    else:
        markdown = _normalize_markdown(await _run_local(_LOCAL_PARSERS[ext], data, ext))

    if not markdown:
        raise ParseError(_EMPTY_MESSAGE)
    return ParsedDocument(markdown=markdown, method=method)
